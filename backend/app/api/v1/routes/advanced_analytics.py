import csv, io
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse, Response
from sqlalchemy.orm import Session
from backend.app.core.security import require_business_owner
from backend.app.db.database import get_db
from backend.app.models.user import User
from backend.app.services.advanced_analytics_service import AdvancedAnalyticsService
from backend.app.schemas.advanced_analytics import AdvancedAnalyticsResponse

router=APIRouter(prefix='/owner/analytics', tags=['Advanced Analytics'])

def report(user: User, db: Session, days: int):
    return AdvancedAnalyticsService.report(db, user.business_id, days)

@router.get('', response_model=AdvancedAnalyticsResponse)
def get_report(days:int=Query(30,ge=1,le=365), user:User=Depends(require_business_owner), db:Session=Depends(get_db)):
    return report(user,db,days)

def _csv_bytes(data:dict):
    out=io.StringIO(); w=csv.writer(out)
    w.writerow(['Metric','Value'])
    for k in ['period_days','total_events','positive_reviews','negative_reviews','positive_ratio','negative_ratio','complaints','google_handoffs','google_handoff_rate','ai_usage','fallback_usage']:
        w.writerow([k,data[k]])
    w.writerow([]); w.writerow(['Rating','Count'])
    for k,v in data['rating_counts'].items(): w.writerow([k,v])
    w.writerow([]); w.writerow(['Source','Count'])
    for k,v in data['source_counts'].items(): w.writerow([k,v])
    return out.getvalue().encode('utf-8-sig')

@router.get('/export.csv')
def export_csv(days:int=Query(30,ge=1,le=365), user:User=Depends(require_business_owner), db:Session=Depends(get_db)):
    data=report(user,db,days)
    return StreamingResponse(io.BytesIO(_csv_bytes(data)), media_type='text/csv', headers={'Content-Disposition': f'attachment; filename=reviewagentai-{data["business_slug"]}-{days}d.csv'})

@router.get('/export.xlsx')
def export_xlsx(days:int=Query(30,ge=1,le=365), user:User=Depends(require_business_owner), db:Session=Depends(get_db)):
    try:
        from openpyxl import Workbook
    except ImportError: raise HTTPException(503,'Excel export dependency is not installed')
    data=report(user,db,days); wb=Workbook(); ws=wb.active; ws.title='Summary'
    for row in [('Metric','Value'), *[(k,data[k]) for k in ['period_days','total_events','positive_reviews','negative_reviews','positive_ratio','negative_ratio','complaints','google_handoffs','google_handoff_rate','ai_usage','fallback_usage']]]: ws.append(row)
    r=wb.create_sheet('Ratings'); r.append(['Rating','Count']); [r.append([k,v]) for k,v in data['rating_counts'].items()]
    s=wb.create_sheet('Sources'); s.append(['Source','Count']); [s.append([k,v]) for k,v in data['source_counts'].items()]
    t=wb.create_sheet('Trend'); t.append(['Period','Events']); [t.append([x['period'],x['events']]) for x in data['monthly_trend']]
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename=reviewagentai-{data["business_slug"]}-{days}d.xlsx'})

@router.get('/export.pdf')
def export_pdf(days:int=Query(30,ge=1,le=365), user:User=Depends(require_business_owner), db:Session=Depends(get_db)):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen.canvas import Canvas
    except ImportError: raise HTTPException(503,'PDF export dependency is not installed')
    data=report(user,db,days); buf=io.BytesIO(); c=Canvas(buf,pagesize=A4); _,h=A4; y=h-50
    c.setFont('Helvetica-Bold',16); c.drawString(45,y,f'ReviewAgentAI Analytics — {data["business_slug"]}'); y-=30
    c.setFont('Helvetica',10)
    rows=[('Period (days)',days),('Total events',data['total_events']),('Positive reviews',data['positive_reviews']),('Negative reviews',data['negative_reviews']),('Positive ratio',f'{data["positive_ratio"]}%'),('Complaints',data['complaints']),('Google handoff rate',f'{data["google_handoff_rate"]}%'),('AI usage',data['ai_usage']),('Fallback usage',data['fallback_usage'])]
    for k,v in rows: c.drawString(55,y,f'{k}: {v}'); y-=18
    y-=8; c.setFont('Helvetica-Bold',11); c.drawString(55,y,'Ratings'); y-=18; c.setFont('Helvetica',10)
    for k,v in data['rating_counts'].items(): c.drawString(65,y,f'{k} stars: {v}'); y-=16
    y-=8; c.setFont('Helvetica-Bold',11); c.drawString(55,y,'Sources'); y-=18; c.setFont('Helvetica',10)
    for k,v in data['source_counts'].items(): c.drawString(65,y,f'{k}: {v}'); y-=16
    c.save(); buf.seek(0)
    return StreamingResponse(buf, media_type='application/pdf', headers={'Content-Disposition': f'attachment; filename=reviewagentai-{data["business_slug"]}-{days}d.pdf'})
